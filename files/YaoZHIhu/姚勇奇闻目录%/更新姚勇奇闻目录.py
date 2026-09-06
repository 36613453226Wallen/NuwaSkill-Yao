#!/usr/bin/env python3
"""读取奇闻目录.json并生成紧凑的姚勇奇闻目录DOCX。"""

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "奇闻目录.json"
OUTPUT_PATH = ROOT / "姚勇奇闻目录.docx"


def set_font(run, name="宋体", size=9, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def compact(paragraph, align=None):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if align is not None:
        paragraph.alignment = align


def document_stats(path):
    if not path.exists():
        return None
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    return {
        "words": len(re.findall(r"[\u3400-\u9fff]", text))
        + len(re.findall(r"[A-Za-z]+(?:[-'][A-Za-z]+)*", text)),
        "headings": sum(p.style.name.startswith("Heading") for p in doc.paragraphs),
        "paragraphs": sum(bool(p.text.strip()) for p in doc.paragraphs),
    }


def generate():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data["records"]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data["title"])
    set_font(r, "黑体", 18, True, (88, 63, 112))

    p = doc.add_paragraph(
        f"登记 {len(records)} 篇｜更新 {data['updated_at']}｜口令：“{data['trigger']}”"
    )
    compact(p, WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        set_font(run, size=8, color=(95, 99, 104))
    p.paragraph_format.space_after = Pt(6)

    headers = ["序号", "DOCX 全名", "时间", "统计与来源"]
    widths = [0.55, 2.75, 1.65, 2.45]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "E8DDF0")
        compact(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            set_font(run, "黑体", 9, True, (88, 63, 112))

    for row_index, record in enumerate(records, start=1):
        stats = document_stats(ROOT / record["docx"])
        stats_text = (
            f"约{stats['words']:,}字｜{stats['headings']}个标题｜"
            f"{stats['paragraphs']}段｜{record['status']}"
            if stats
            else f"{record['status']}｜DOCX尚未生成"
        )
        values = [
            str(record["id"]),
            record["docx"],
            f"发布：{record['published']}\n编辑：{record['edited']}",
            stats_text,
        ]
        row = table.add_row()
        for i, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[i]
            cell.width = Inches(width)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                shade(cell, "FAF7FC")
            compact(
                cell.paragraphs[0],
                WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT,
            )
            for run in cell.paragraphs[0].runs:
                set_font(run, size=8.5)
        p = row.cells[3].add_paragraph("原文：" + record["source_url"])
        compact(p)
        for run in p.runs:
            set_font(run, size=7.5, color=(88, 63, 112))

    doc.core_properties.title = data["title"]
    doc.save(OUTPUT_PATH)
    print(f"已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    generate()

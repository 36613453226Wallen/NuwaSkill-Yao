#!/usr/bin/env python3
"""读取文章目录.json并生成紧凑、美观的姚勇文章目录DOCX。"""

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "文章目录.json"
OUTPUT_PATH = ROOT / "姚勇文章目录.docx"


def set_font(run, name="宋体", size=9, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def compact(paragraph, align=None):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if align is not None:
        paragraph.alignment = align


def document_stats(path):
    if not path.exists():
        return None
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    text = "\n".join(texts)
    cjk = len(re.findall(r"[\u3400-\u9fff]", text))
    latin_words = len(re.findall(r"[A-Za-z]+(?:[-'][A-Za-z]+)*", text))
    headings = sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading"))
    return {
        "word_count": cjk + latin_words,
        "paragraphs": sum(1 for value in texts if value.strip()),
        "headings": headings,
    }


def generate():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data["records"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(data["title"])
    set_font(r, "黑体", 18, True, (31, 78, 121))

    completed = sum(1 for record in records if (ROOT / record["docx"]).exists())
    p = doc.add_paragraph(
        f"登记 {len(records)} 篇｜已生成 {completed} 篇｜更新 {data['updated_at']}｜口令：“{data['trigger']}”"
    )
    compact(p, WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        set_font(run, "宋体", 8, color=(95, 99, 104))
    p.paragraph_format.space_after = Pt(6)

    headers = ["序号", "DOCX 全名", "时间", "内容统计与状态"]
    widths = [0.55, 2.75, 1.65, 2.45]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "D9EAF7")
        compact(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            set_font(run, "黑体", 9, True, (31, 78, 121))

    for row_index, record in enumerate(records, start=1):
        path = ROOT / record["docx"]
        stats = document_stats(path)
        if stats:
            stats_text = (
                f"约{stats['word_count']:,}字｜{stats['headings']}个标题｜"
                f"{stats['paragraphs']}段｜{record['status']}"
            )
        else:
            stats_text = f"{record['status']}｜DOCX尚未生成"

        time_text = f"发布：{record['published']}\n编辑：{record['edited']}"
        docx_name = record["docx"]
        if record.get("extra_docx"):
            docx_name = f"{record['docx']}\n{record['extra_docx']}"
        values = [str(record["id"]), docx_name, time_text, stats_text]
        row = table.add_row()
        for i, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[i]
            cell.width = Inches(width)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                shade(cell, "F6F8FA")
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            compact(cell.paragraphs[0], align)
            for run in cell.paragraphs[0].runs:
                set_font(run, "宋体", 8.5)

        if record.get("source_url"):
            source = row.cells[3].add_paragraph("原文：" + record["source_url"])
            compact(source)
            for run in source.runs:
                set_font(run, "宋体", 7.5, color=(46, 116, 181))

    doc.add_paragraph()
    p = doc.add_paragraph(
        "维护规则：新增PDF后按序号生成DOCX；首次发布时间、后续编辑时间与快照时间分别记录。"
        "内容统计只计算DOCX正文中的汉字与英文词，不把标点计入。"
    )
    compact(p)
    for run in p.runs:
        set_font(run, "宋体", 8, color=(89, 89, 89))

    doc.core_properties.title = data["title"]
    doc.save(OUTPUT_PATH)
    print(f"已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    generate()

#!/usr/bin/env python3
"""Generate a formatted Agent / Yao Yong dual-perspective DOCX."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, RGBColor

from docx_style import (
    AGENT_BLUE,
    ASK_GREEN,
    INK,
    NAVY,
    SLATE,
    YAO_BROWN,
    add_bottom_border,
    compact,
    shade_cell,
    set_cell_margins,
    set_run_font,
)


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "双版本记录.json"

SPEAKER_STYLE = {
    "提问人": {"fill": "E8F5EE", "color": ASK_GREEN, "accent": "2E6B4A"},
    "Agent": {"fill": "E8F1F8", "color": AGENT_BLUE, "accent": "1B5E8A"},
    "姚勇视角": {"fill": "F8F0E6", "color": YAO_BROWN, "accent": "8A4B12"},
}


def status_of(record: dict) -> tuple[str, str]:
    speakers = record.get("speakers") or {}
    question = (speakers.get("提问人") or "").strip()
    agent = (speakers.get("Agent") or "").strip()
    yao = (speakers.get("姚勇视角") or "").strip()
    if not question:
        return "没有问题", "缺少可摘录的提问，不编造问题。"
    n = int(bool(agent)) + int(bool(yao))
    if n == 0:
        return "有问题、没有回答", "仅有提问，两份回答都缺。"
    if n == 1:
        missing = "姚勇视角" if agent else "Agent"
        return "只有一份回答", f"已有提问和一份回答，缺「{missing}」。"
    return "双份回答", "提问、Agent、姚勇视角均已齐。"


def add_heading_run(paragraph, text, size=18):
    run = paragraph.add_run(text)
    set_run_font(run, "黑体", size, True, NAVY)
    return run


def add_meta_table(doc, record, folder_brief, status, status_note):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rows = [
        ("日期", record.get("date") or "约/未知"),
        ("完成状态", f"{status}｜{status_note}"),
        ("发言人标记", "表头用「发言人：」，后接摘录正文"),
        ("文件夹", folder_brief),
    ]
    widths = (Inches(1.35), Inches(5.55))
    for i, (label, value) in enumerate(rows):
        left, right = table.rows[i].cells
        left.text = ""
        right.text = ""
        left.width, right.width = widths
        shade_cell(left, "1F4E79")
        shade_cell(right, "F4F7FA" if i % 2 == 0 else "FFFFFF")
        set_cell_margins(left, top=60, bottom=60, left=80, right=80)
        set_cell_margins(right, top=60, bottom=60, left=100, right=80)
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        compact(lp, after=0, line=1.08)
        compact(rp, after=0, line=1.08)
        lr = lp.add_run(label)
        rr = rp.add_run(value)
        set_run_font(lr, "黑体", 9, True, RGBColor(0xFF, 0xFF, 0xFF))
        set_run_font(rr, "宋体", 9, False, INK)
        for cell in (left, right):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_speaker_block(doc, speaker, body):
    style = SPEAKER_STYLE.get(speaker, SPEAKER_STYLE["Agent"])
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.9)
    shade_cell(cell, style["fill"])
    set_cell_margins(cell, top=80, bottom=120, left=140, right=140)

    header = cell.paragraphs[0]
    compact(header, before=2, after=8, line=1.0)
    add_bottom_border(header, style["accent"], "8")
    run = header.add_run(f"{speaker}：")
    set_run_font(run, "黑体", 13, True, style["color"])

    text = (body or "").strip()
    if not text:
        empty = cell.add_paragraph()
        compact(empty, after=2, line=1.15)
        er = empty.add_run("（本条缺失，未编造。）")
        set_run_font(er, "楷体", 10.5, False, SLATE, italic=True)
        return

    paragraphs = [part.strip() for part in text.split("\n") if part.strip()]
    for index, part in enumerate(paragraphs):
        p = cell.add_paragraph()
        compact(p, after=6 if index < len(paragraphs) - 1 else 2, line=1.2, first_line=22)
        r = p.add_run(part)
        font_name = "楷体" if speaker == "姚勇视角" else "宋体"
        set_run_font(r, font_name, 11, False, INK)


def generate_record(record: dict, folder_brief: str) -> Path:
    status, status_note = status_of(record)
    record["question_status"] = "有问题" if (record.get("speakers") or {}).get("提问人", "").strip() else "没有问题"
    record["answer_status"] = status

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    kicker = doc.add_paragraph()
    compact(kicker, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
    kr = kicker.add_run("问题与视角  ·  Agent 与 Y 双版本")
    set_run_font(kr, "黑体", 10, True, NAVY)

    title = doc.add_paragraph()
    compact(title, before=0, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_bottom_border(title, "1F4E79", "16")
    add_heading_run(title, record["title"], 18)

    subtitle = doc.add_paragraph()
    compact(subtitle, before=4, after=10, line=1.1)
    sr = subtitle.add_run("同一问题先留 Agent 判断，再留姚勇公开框架下的一版。发言人用表头加冒号标出。")
    set_run_font(sr, "宋体", 10, False, SLATE)

    add_meta_table(doc, record, folder_brief, status, status_note)
    doc.add_paragraph()

    speakers = record.get("speakers") or {}
    for name in ("提问人", "Agent", "姚勇视角"):
        add_speaker_block(doc, name, speakers.get(name, ""))
        spacer = doc.add_paragraph()
        compact(spacer, after=8, line=1.0)

    note = doc.add_paragraph()
    compact(note, before=2, after=0, line=1.15)
    nr = note.add_run(
        "说明：姚勇视角基于其公开文章与访谈的模拟，不代表本人。文件夹「问题与视角-agent与Y双版本」"
        "用于保存这类双份摘录，并集中登记口令总表《*超级指令》。"
    )
    set_run_font(nr, "宋体", 9, False, SLATE)

    output = ROOT / record["docx"]
    doc.core_properties.title = record["title"]
    doc.core_properties.subject = "Agent与姚勇双视角摘录"
    doc.save(output)
    return output


def describe_folder(data: dict) -> str:
    return data.get("folder_brief") or data.get("folder")


def generate(record_id: int | None = None) -> list[Path]:
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data["records"]
    if record_id is None:
        chosen = records
    else:
        chosen = [item for item in records if item["id"] == record_id]
        if not chosen:
            raise SystemExit(f"没有 id={record_id} 的双版本记录")

    paths = []
    for record in chosen:
        path = generate_record(record, describe_folder(data))
        status, note = status_of(record)
        print(f"状态：{status}")
        print(f"说明：{note}")
        print(f"文件夹：{data['folder']}｜{describe_folder(data)}")
        print(f"已生成：{path}")
        paths.append(path)

    DATA_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return paths


if __name__ == "__main__":
    target = int(sys.argv[1]) if len(sys.argv) > 1 else None
    generate(target)
